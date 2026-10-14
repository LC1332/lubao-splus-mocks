#!/usr/bin/env python3
"""Generate validated floor-and-wall block matching questions and a DOCX."""
from __future__ import annotations

import argparse
from io import BytesIO
import json
from pathlib import Path

from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt

from block_matching_generator import make_questions
from block_matching_render import draw_floor, draw_wall, draw_option

ROOT = Path(__file__).resolve().parents[1]


def explanation(option: dict) -> str:
    if option["is_answer"]:
        return "正确"
    errors = []
    if not option["count_ok"]:
        errors.append("数量错误")
    if not option["floor_ok"]:
        errors.append("地板不符")
    if not option["wall_ok"]:
        errors.append("墙洞轮廓不符")
    return "、".join(errors)


def font(size: int):
    for path in ("/System/Library/Fonts/PingFang.ttc", "/System/Library/Fonts/STHeiti Light.ttc",
                 "/usr/share/fonts/truetype/noto/NotoSansCJK-Regular.ttc"):
        if Path(path).exists():
            return ImageFont.truetype(path, size=size)
    return ImageFont.load_default()


def add_cjk(paragraph, text: str, size_pt=11, bold=False, color="#28343D") -> None:
    dpi = 300
    fnt = font(round(size_pt * dpi / 72))
    pad = round(dpi / 45)
    scratch = Image.new("RGBA", (8, 8), (255, 255, 255, 0))
    draw = ImageDraw.Draw(scratch)
    box = draw.textbbox((0, 0), text, font=fnt)
    w, h = max(1, box[2] - box[0] + pad * 2), max(1, box[3] - box[1] + pad * 2)
    image = Image.new("RGBA", (w, h), (255, 255, 255, 0))
    ImageDraw.Draw(image).text((pad - box[0], pad - box[1]), text, font=fnt, fill=color,
                               stroke_width=(1 if bold else 0), stroke_fill=color)
    stream = BytesIO()
    image.save(stream, format="PNG", dpi=(dpi, dpi))
    stream.seek(0)
    run = paragraph.add_run()
    pic = run.add_picture(stream, width=Inches(w / dpi), height=Inches(h / dpi))
    pic._inline.docPr.set("descr", text)


def set_cell_margins(cell, top=55, start=70, bottom=45, end=70):
    tc_pr = cell._tc.get_or_add_tcPr()
    margins = tc_pr.first_child_found_in("w:tcMar")
    if margins is None:
        margins = OxmlElement("w:tcMar")
        tc_pr.append(margins)
    for edge, val in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = margins.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            margins.append(node)
        node.set(qn("w:w"), str(val))
        node.set(qn("w:type"), "dxa")


def no_borders(table):
    pr = table._tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = OxmlElement(f"w:{edge}")
        tag.set(qn("w:val"), "nil")
        borders.append(tag)
    pr.append(borders)


def image_alt(run, text: str):
    pic = run.add_picture(str(text))
    pic._inline.docPr.set("descr", Path(text).stem)
    return pic


def make_docx(questions: list[dict], assets: Path, out_path: Path, wall_rule: str = "pass") -> None:
    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Inches(8.27), Inches(11.69)
    sec.top_margin, sec.bottom_margin = Inches(.42), Inches(.42)
    sec.left_margin, sec.right_margin = Inches(.58), Inches(.58)
    sec.header_distance, sec.footer_distance = Inches(.2), Inches(.2)

    for qi, q in enumerate(questions):
        if qi:
            doc.add_page_break()
        if qi == 0:
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_after = Pt(3)
            add_cjk(p, "地板与墙洞匹配立体积木练习", 20, True, "#000000")
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_after = Pt(5)
            add_cjk(p, "姓名：____________    日期：____________", 10)
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(2)
        p.paragraph_format.space_after = Pt(3)
        add_cjk(p, f"第 {q['number']} 题", 14, True)
        stem = ("　选出由 6 块积木组成、占满地板且正面轮廓与墙洞完全一致的组合。不能旋转，左右按图对应。"
                if wall_rule == "exact" else
                "　选出由 6 块积木组成、占满地板并能穿过墙洞的组合。不能旋转，左右按图对应。")
        add_cjk(p, stem, 10)

        targets = doc.add_table(rows=2, cols=2)
        targets.autofit = False
        targets.columns[0].width = Inches(3.5)
        targets.columns[1].width = Inches(3.5)
        no_borders(targets)
        for col, (label, fname) in enumerate((("地板俯视图", f"q{q['number']:02d}_floor.png"),
                                               ("墙洞正面图", f"q{q['number']:02d}_wall.png"))):
            head = targets.cell(0, col)
            head.width = Inches(3.5)
            head.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            p = head.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_after = Pt(0)
            add_cjk(p, label, 10, True)
            cell = targets.cell(1, col)
            cell.width = Inches(3.5)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell, 0, 0, 0, 0)
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_after = Pt(1)
            image_alt(p.add_run(), str(assets / fname)).width = Inches(2.62)

        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(1)
        add_cjk(p, "积木选项", 11, True)
        options = doc.add_table(rows=2, cols=2)
        options.autofit = False
        options.columns[0].width = Inches(3.55)
        options.columns[1].width = Inches(3.55)
        no_borders(options)
        for idx, option in enumerate(q["options"]):
            cell = options.cell(idx // 2, idx % 2)
            cell.width = Inches(3.55)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell, 15, 55, 20, 55)
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_after = Pt(0)
            add_cjk(p, option["letter"], 13, True)
            p = cell.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_after = Pt(0)
            pic = image_alt(p.add_run(), str(assets / f"q{q['number']:02d}_option_{option['letter']}.png"))
            pic.width = Inches(2.9)
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        p.paragraph_format.space_before = Pt(1)
        p.paragraph_format.space_after = Pt(0)
        add_cjk(p, "答案：（　　　）", 12)

    # Answers start on a separate page and use a compact, legible four-question rhythm.
    doc.add_page_break()
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(4)
    add_cjk(p, "参考答案与逐项判断", 19, True, "#000000")
    for qi, q in enumerate(questions):
        if qi and qi % 4 == 0:
            doc.add_page_break()
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_after = Pt(4)
            add_cjk(p, "参考答案与逐项判断", 17, True, "#000000")
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(5)
        p.paragraph_format.space_after = Pt(1)
        add_cjk(p, f"第 {q['number']} 题　正确答案：{q['answer']}", 12, True)
        for option in q["options"]:
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Inches(.18)
            p.paragraph_format.space_after = Pt(0)
            count_text = f"{option['cube_count']} 块"
            floor_text = "地板正确" if option["floor_ok"] else "地板不符"
            wall_text = "能通过" if option["wall_ok"] else "不能通过"
            reason = explanation(option)
            add_cjk(p, f"{option['letter']}：{count_text}｜{floor_text}｜{wall_text}｜{reason}", 9.5)

    out_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(out_path)


def build(args) -> Path:
    out = args.output_dir.expanduser().resolve()
    assets = out / "立体积木匹配素材"
    assets.mkdir(parents=True, exist_ok=True)
    questions = make_questions(args.count, args.seed, args.wall_rule)
    for q in questions:
        draw_floor(q["target_floor"], assets / f"q{q['number']:02d}_floor.png")
        draw_wall(q["target_wall"], assets / f"q{q['number']:02d}_wall.png")
        for option in q["options"]:
            draw_option(option["matrix"], option["cube_colors"], assets / f"q{q['number']:02d}_option_{option['letter']}.png")
            option["height_matrix"] = [list(row) for row in option.pop("matrix")]
            option["cube_colors"] = [{"coordinate": list(coord), "base_color": color}
                                      for coord, color in sorted(option["cube_colors"].items())]
            option["image"] = f"q{q['number']:02d}_option_{option['letter']}.png"
    payload = {"seed": args.seed, "configuration": {"count": args.count, "target_count": 6,
                "wall_rule": args.wall_rule, "questions_per_page": 1,
                "projection": {"type": "parallel_isometric", "x_unit": 47, "depth_x": 34, "depth_y_ratio": 0.78}},
                "questions": questions,
                "answers": [{"question": q["number"], "answer": q["answer"],
                             "options": [{"letter": o["letter"], "cube_count": o["cube_count"],
                                          "floor_ok": o["floor_ok"], "wall_ok": o["wall_ok"],
                                          "truth": o["truth"], "reason": explanation(o)}
                                         for o in q["options"]]} for q in questions]}
    (assets / "题目参数与答案.json").write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")
    docx = out / "大班数学立体积木匹配练习.docx"
    make_docx(questions, assets, docx, args.wall_rule)
    return docx


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--count", type=int, default=12, help="题目数量，默认 12")
    parser.add_argument("--seed", type=int, default=20261014, help="复现随机种子")
    parser.add_argument("--output-dir", type=Path, default=ROOT / "generated", help="输出目录")
    parser.add_argument("--wall-rule", choices=("pass", "exact"), default="pass", help="墙洞判定规则")
    args = parser.parse_args()
    if not 1 <= args.count <= 12:
        parser.error("--count 必须为 1 至 12")
    try:
        result = build(args)
    except RuntimeError as exc:
        parser.error(str(exc))
    print(f"已生成：{result}")
    print(f"题目：{args.count} 道；随机种子：{args.seed}；每页 1 题；墙洞规则：{args.wall_rule}")


if __name__ == "__main__":
    main()
