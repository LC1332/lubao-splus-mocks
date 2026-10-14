#!/usr/bin/env python3
"""Generate the 12-question square-counting worksheet and answer key."""
from __future__ import annotations

import argparse
from io import BytesIO
import json
from pathlib import Path

from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt

from square_count_logic import make_questions, validate
from square_count_render import draw_grid

ROOT = Path(__file__).resolve().parents[1]
TITLE = "在组合图形中数正方形"
STEM = "请问下面图中有多少个正方形？把答案填在括号里。"


def set_run_font(run, size: float, bold: bool = False) -> None:
    run.font.name = "Arial"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    run.font.size = Pt(size)
    run.bold = bold


def cjk_font(size_px: int):
    for candidate in ("/System/Library/Fonts/PingFang.ttc",
                      "/System/Library/Fonts/STHeiti Light.ttc",
                      "/usr/share/fonts/truetype/noto/NotoSansCJK-Regular.ttc"):
        if Path(candidate).exists():
            return ImageFont.truetype(candidate, size=size_px)
    return ImageFont.load_default()


def add_text(paragraph, text: str, size: float, bold: bool = False):
    # Render text with a known CJK font so headless Word renderers and school
    # computers without the same font fallback still show Chinese correctly.
    dpi = 300
    font = cjk_font(round(size * dpi / 72))
    pad = round(dpi / 36)
    scratch = Image.new("RGBA", (8, 8), (255, 255, 255, 0))
    box = ImageDraw.Draw(scratch).textbbox((0, 0), text, font=font,
                                           stroke_width=1 if bold else 0)
    width = max(1, box[2] - box[0] + 2 * pad)
    height = max(1, box[3] - box[1] + 2 * pad)
    raster = Image.new("RGBA", (width, height), (255, 255, 255, 0))
    ImageDraw.Draw(raster).text((pad - box[0], pad - box[1]), text, font=font,
                                fill="#111111", stroke_width=1 if bold else 0,
                                stroke_fill="#111111")
    stream = BytesIO()
    raster.save(stream, format="PNG", dpi=(dpi, dpi))
    stream.seek(0)
    run = paragraph.add_run()
    picture = run.add_picture(stream, width=Inches(width / dpi), height=Inches(height / dpi))
    picture._inline.docPr.set("descr", text)
    return run


def no_table_borders(table) -> None:
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        node = OxmlElement(f"w:{edge}")
        node.set(qn("w:val"), "nil")
        borders.append(node)
    table._tbl.tblPr.append(borders)


def prevent_row_split(row) -> None:
    row_properties = row._tr.get_or_add_trPr()
    row_properties.append(OxmlElement("w:cantSplit"))


def make_docx(questions: list[dict], assets: Path, output: Path) -> None:
    doc = Document()
    section = doc.sections[0]
    section.page_width, section.page_height = Inches(8.27), Inches(11.69)
    section.top_margin, section.bottom_margin = Inches(.55), Inches(.55)
    section.left_margin, section.right_margin = Inches(.72), Inches(.72)
    section.header_distance = section.footer_distance = Inches(.25)

    for index, question in enumerate(questions):
        if index % 2 == 0 and index:
            doc.add_page_break()
        if index == 0:
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_after = Pt(5)
            add_text(p, TITLE, 20, True)
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_after = Pt(14)
            add_text(p, "姓名：________________　日期：________________", 11)

        # A one-row table keeps each question, picture and answer space together.
        table = doc.add_table(rows=1, cols=1)
        no_table_borders(table)
        row = table.rows[0]
        prevent_row_split(row)
        cell = row.cells[0]
        cell.width = Inches(6.75)
        p = cell.paragraphs[0]
        p.paragraph_format.space_before = Pt(4)
        p.paragraph_format.space_after = Pt(4)
        add_text(p, f"{question['number']}．{STEM}", 12, True)
        p = cell.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(2)
        picture = p.add_run().add_picture(str(assets / f"q{question['number']:02d}.png"))
        picture.width = Inches(2.35)
        picture._inline.docPr.set("descr", f"第{question['number']}题组合图形")
        p = cell.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        p.paragraph_format.space_after = Pt(7)
        add_text(p, "（　　　　）个", 13)

    doc.add_page_break()
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(12)
    add_text(p, "参考答案", 20, True)
    for q in questions:
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(5)
        p.paragraph_format.space_after = Pt(2)
        add_text(p, f"第 {q['number']} 题：共 {q['total']} 个正方形。", 12, True)
        count_text = "；".join(f"{k}×{k}：{q['counts'][str(k)]} 个" for k in range(1, 6))
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Inches(.2)
        p.paragraph_format.space_after = Pt(0)
        add_text(p, f"{count_text}。", 10.5)
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Inches(.2)
        p.paragraph_format.space_after = Pt(3)
        add_text(p, f"合计：{' + '.join(str(q['counts'][str(k)]) for k in range(1, 6))} = {q['total']}。", 10.5)

    output.parent.mkdir(parents=True, exist_ok=True)
    doc.save(output)


def build(count: int, seed: int, output_dir: Path) -> tuple[Path, Path]:
    output_dir = output_dir.expanduser().resolve()
    assets = output_dir / "组合图形数正方形素材"
    questions = make_questions(count, seed)
    for q in questions:
        # Recheck the exact data immediately before drawing and serialization.
        report = validate(tuple(tuple(row) for row in q["matrix"]))
        if not report["valid"] or report["total"] != q["total"]:
            raise RuntimeError(f"第 {q['number']} 题最终校验失败。")
        q["checks"] = report["checks"]
        draw_grid(q["matrix"], assets / f"q{q['number']:02d}.png")
    payload = {
        "generation_version": "1.0",
        "seed": seed,
        "configuration": {"count": count, "rows": 5, "columns": 6,
                          "questions_per_page": 2, "max_attempts": 200000,
                          "cell_color": "#FFD84D", "grid_color": "#1A1A1A"},
        "questions": questions,
    }
    data_path = assets / "题目数据与答案.json"
    data_path.parent.mkdir(parents=True, exist_ok=True)
    data_path.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")
    docx_path = output_dir / "大班数学组合图形数正方形练习.docx"
    make_docx(questions, assets, docx_path)
    return docx_path, data_path


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--count", type=int, default=12, help="题目数，默认 12（最多 12）")
    parser.add_argument("--seed", type=int, default=20261014, help="复现随机种子")
    parser.add_argument("--output-dir", type=Path, default=ROOT / "generated", help="输出目录")
    args = parser.parse_args()
    if not 1 <= args.count <= 12:
        parser.error("--count 必须为 1 至 12")
    try:
        docx, data = build(args.count, args.seed, args.output_dir)
    except RuntimeError as exc:
        parser.error(str(exc))
    print(f"已生成：{docx}")
    print(f"题目数据：{data}")
    print(f"题数：{args.count}；随机种子：{args.seed}；每页 2 题")


if __name__ == "__main__":
    main()
